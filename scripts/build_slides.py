from __future__ import annotations
from pathlib import Path
import os
import re
import subprocess
import sys
import tempfile
import time
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from common import ROOT, course, summary_parts, section_text

META = course()
VERSION = META["version"]
BOOK = ROOT / "book/zh"
OUT = ROOT / "slides"
OUT.mkdir(exist_ok=True)


def normalize_ooxml_zip(path: Path):
    """Rewrite OOXML with stable member order and SOURCE_DATE_EPOCH timestamps.

    python-pptx delegates ZIP timestamps to zipfile, which otherwise records the
    wall clock.  The XML/package payload is deterministic; normalizing the ZIP
    envelope makes clean builds byte-identical without changing presentation
    semantics.
    """
    epoch = int(os.environ.get("SOURCE_DATE_EPOCH", "1789084800"))
    dt = time.gmtime(max(epoch, 315532800))[:6]  # ZIP timestamps cannot predate 1980.
    with zipfile.ZipFile(path, "r") as zin:
        members = [(i.filename, zin.read(i.filename), i.compress_type) for i in zin.infolist()]
    fd, tmp = tempfile.mkstemp(prefix=path.stem + ".", suffix=".pptx", dir=path.parent)
    os.close(fd)
    try:
        with zipfile.ZipFile(tmp, "w") as zout:
            for name, data, ctype in sorted(members, key=lambda x: x[0]):
                zi = zipfile.ZipInfo(name, dt)
                zi.create_system = 3
                zi.external_attr = 0o100644 << 16
                zi.compress_type = ctype
                zout.writestr(zi, data, compress_type=ctype, compresslevel=9 if ctype == zipfile.ZIP_DEFLATED else None)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


DIAG = ROOT / "book/assets/diagrams"
subprocess.run([sys.executable, str(ROOT / "scripts/build_diagrams.py")], check=True)

FONT = "Noto Sans CJK SC"
NAVY = RGBColor(31, 78, 121)
INK = RGBColor(32, 39, 46)
MUTED = RGBColor(88, 96, 105)
LIGHT = RGBColor(246, 248, 250)
GREEN = RGBColor(52, 122, 88)
ORANGE = RGBColor(181, 92, 47)
WHITE = RGBColor(255, 255, 255)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bg(slide, c=WHITE):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = c


def tx(slide, text, x, y, w, h, size=20, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return box


def title(slide, t, sub=None, no=None):
    if no:
        tx(slide, no, 0.68, 0.35, 1.2, 0.4, 16, True, NAVY)
    tx(slide, t, 0.72, 0.78, 11.9, 0.95, 28, True, INK)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.75), Inches(11.9), Inches(0.035))
    line.fill.solid()
    line.fill.fore_color.rgb = NAVY
    line.line.fill.background()
    if sub:
        tx(slide, sub, 0.75, 1.88, 11.7, 0.58, 14, False, MUTED)


def bullets(slide, items, x, y, w, h, size=15, color=INK, max_items=6):
    items = [i.strip() for i in items if i and i.strip()]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, it in enumerate(items[:max_items]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + it
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def picture_contain(slide, path, x, y, w, h):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), width=Inches(pw), height=Inches(ph)
    )


def clean_md(s: str) -> str:
    s = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", s)
    s = re.sub(r"[`*_]", "", s)
    return re.sub(r"\s+", " ", s).strip()


def parse_chapter(path: Path):
    s = path.read_text(encoding="utf-8")
    m = re.search(r"^# (.+)", s, re.M)
    ttl = m.group(1) if m else path.stem
    thesis = re.search(r"> \*\*本章核心判断\*\*：(.+)", s)
    thesis = clean_md(thesis.group(1)) if thesis else ""
    inv = re.search(r"> \*\*Invariant\*\*[：:]\s*(.+)", s)
    invariant = clean_md(inv.group(1)) if inv else ""
    concepts = [
        clean_md(x) for x in re.findall(r"^###\s+(.+)$", section_text(s, "核心概念与系统直觉", "原理与理论基础"), re.M)
    ]
    principles = ([f"Invariant: {invariant}"] if invariant else []) + concepts[:4]
    steps = [
        clean_md(x)
        for x in re.findall(
            r"^\*\*Step\s+\d+\s+[—-]\s+(.+?)。?\*\*", section_text(s, "关键机制与执行流程", "从原理到实现"), re.M
        )
    ]
    failure_sec = section_text(s, "故障模型、失败模式与排错", "性能、可靠性与工程化")
    failures = [clean_md(x) for x in re.findall(r"^-\s+\*\*(.+?)\*\*[:：]", failure_sec, re.M)]
    if not failures:
        failures = [clean_md(x) for x in re.findall(r"^###\s+(.+)$", failure_sec, re.M)]
    projects = []
    for row in section_text(s, "主流系统实现对照与源码阅读入口", "设计方案与方法对比").splitlines():
        if row.startswith("|") and "---" not in row and "项目" not in row:
            cols = [clean_md(x) for x in row.strip("|").split("|")]
            if cols and cols[0]:
                projects.append(cols[0])
    labsec = section_text(s, "可复现实验", "工程场景与系统设计")
    labs = [clean_md(x) for x in re.findall(r"^###\s+(Lab\s+[^\n]+)$", labsec, re.M)]
    arch = DIAG / f"{path.stem}-architecture.png"
    flow = DIAG / f"{path.stem}-flow.png"
    return ttl, thesis, principles, steps, failures, projects, labs, arch, flow


# Cover
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, LIGHT)
tx(s, META["title"], 0.8, 1.15, 11.8, 0.9, 45, True, NAVY)
tx(s, META["book_tagline"], 0.85, 2.2, 11.5, 0.6, 24, False, INK)
tx(s, f"{VERSION} · 技术教材 / 工程实践 / 源码解析 / 研究型教程", 0.85, 3.1, 11.6, 0.55, 18, True, GREEN)
tx(
    s,
    "问题 -> 直觉 -> 原理 -> 实现 -> 源码 -> 实验 -> 故障 -> 工程 -> 研究",
    0.9,
    5.0,
    11.4,
    0.6,
    18,
    False,
    MUTED,
    PP_ALIGN.CENTER,
)

s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
title(s, "全书学习闭环", "40 章递进：每章建立不变量，并用正常路径与故障路径验证")
mapdiag = DIAG / "01-foundation-architecture.png"
if mapdiag.exists():
    picture_contain(s, mapdiag, 0.9, 2.55, 11.5, 3.75)

idx = 0
for part in summary_parts():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, NAVY)
    tx(s, part["title"], 0.8, 2.18, 11.8, 1.1, 31, True, WHITE, PP_ALIGN.CENTER)
    tx(
        s,
        f"{len(part['chapters'])} 章 · 原理 / 实现 / 源码 / 实验 / 故障 / 工程 / 研究",
        1.1,
        3.55,
        11.1,
        0.6,
        17,
        False,
        RGBColor(227, 238, 248),
        PP_ALIGN.CENTER,
    )
    for ch in part["chapters"]:
        idx += 1
        path = BOOK / ch["path"]
        ttl, thesis, principles, steps, failures, projects, labs, arch, flow = parse_chapter(path)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg(s)
        title(s, ttl, thesis[:150], f"{idx:02d}")
        if arch.exists():
            picture_contain(s, arch, 7.05, 2.55, 5.25, 3.55)
        tx(s, "原理与不变量", 0.78, 2.58, 5.9, 0.35, 17, True, NAVY)
        bullets(s, principles[:5] or steps[:5], 0.82, 3.02, 5.8, 2.35, 14, max_items=5)
        tx(s, "典型失败", 0.78, 5.68, 5.9, 0.35, 16, True, ORANGE)
        bullets(s, failures[:3], 0.82, 6.03, 5.8, 0.85, 13, max_items=3)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg(s)
        title(s, ttl + "：实验、源码与工程化", "核心实验本地确定性验证；上游项目按 SOURCE_LOCK 独立复现", f"{idx:02d}")
        if flow.exists():
            picture_contain(s, flow, 8.35, 2.55, 4.1, 2.9)
        tx(s, "核心实验", 0.82, 2.58, 3.35, 0.35, 17, True, GREEN)
        bullets(s, labs[:4], 0.87, 3.0, 3.25, 1.7, 14, max_items=4)
        tx(s, "开源源码对照", 4.38, 2.58, 3.55, 0.35, 17, True, NAVY)
        bullets(s, projects[:5], 4.43, 3.0, 3.5, 2.25, 13, max_items=5)
        tx(s, "验收证据", 0.85, 5.42, 7.25, 0.35, 16, True, INK)
        bullets(
            s,
            ["环境/版本/完整命令", "实际输出与 PASS/FAIL", "关键断点与状态变化", "故障注入后不变量仍成立"],
            0.9,
            5.78,
            7.1,
            1.15,
            13,
            max_items=4,
        )

s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, LIGHT)
tx(s, "完成课程后，应能独立设计 Agent Systems", 0.85, 1.25, 11.6, 0.75, 34, True, NAVY, PP_ALIGN.CENTER)
bullets(
    s,
    [
        "从零实现 Agent loop、Tool Runtime、Checkpoint/Journal 与 verifier",
        "能沿真实源码调用链理解 Agents SDK、ADK、LangGraph、MAF、DeepSeek Harness、Pi、Codex、OpenHands",
        "能把正常路径与故障路径放进同一套可复现实验",
        "能用 trace、journal、checkpoint、external observation 解释系统事实",
        "能从工业约束出发识别 Research Gap 与开放问题",
    ],
    1.75,
    2.75,
    9.8,
    3.2,
    18,
    max_items=5,
)
path = OUT / f"AI-Agent-Systems-Course-{VERSION}.pptx"
prs.save(path)
normalize_ooxml_zip(path)
print("SLIDES_BUILT", len(prs.slides), path)

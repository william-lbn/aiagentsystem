from __future__ import annotations
import re
from pptx import Presentation
from common import ROOT, course, build_cfg, chapter_paths, section_text

cfg = build_cfg()
meta = course()
errors = []


def req(c, m):
    if not c:
        errors.append(m)


for p in chapter_paths():
    s = p.read_text(encoding="utf-8")
    inv = re.search(r"> \*\*Invariant\*\*[：:]\s*(.+)", s)
    concepts = re.findall(r"^###\s+(.+)$", section_text(s, "核心概念与系统直觉", "原理与理论基础"), re.M)
    principles = ([inv.group(1)] if inv else []) + concepts
    projects = []
    for row in section_text(s, "主流系统实现对照与源码阅读入口", "设计方案与方法对比").splitlines():
        if row.startswith("|") and "---" not in row and "项目" not in row:
            cols = [x.strip() for x in row.strip("|").split("|")]
            if cols and cols[0]:
                projects.append(cols[0])
    labs = re.findall(r"^###\s+(Lab\s+[^\n]+)$", section_text(s, "可复现实验", "工程场景与系统设计"), re.M)
    failures = section_text(s, "故障模型、失败模式与排错", "性能、可靠性与工程化")
    req(bool(principles), f"{p.name}: slide parser would have no principles")
    req(bool(projects), f"{p.name}: slide parser would have no upstream projects")
    req(len(labs) >= 2, f"{p.name}: slide parser labs={len(labs)}")
    req(bool(failures.strip()), f"{p.name}: failure section empty")
ppt = ROOT / "slides" / f"AI-Agent-Systems-Course-{meta['version']}.pptx"
req(ppt.exists(), f"missing {ppt}")
if ppt.exists():
    mode = ppt.stat().st_mode & 0o777
    req((mode & 0o444) == 0o444, f"slide artifact mode={mode:o} must be host-readable")
    prs = Presentation(ppt)
    req(len(prs.slides) == cfg["expected_slides"], f"slides={len(prs.slides)} expected {cfg['expected_slides']}")
    alltext = "\n".join(sh.text for sl in prs.slides for sh in sl.shapes if hasattr(sh, "text"))
    req("SOURCE_LOCK 中锁定的官方实现" not in alltext, "slides still contain old generic upstream fallback")
    req("违反不变量时必须进入显式失败/恢复路径" not in alltext, "slides still contain old generic failure fallback")
if errors:
    print("SLIDES_QA_FAILED")
    [print("-", x) for x in errors]
    raise SystemExit(1)
print(f"SLIDES_QA_OK slides={cfg['expected_slides']}")
